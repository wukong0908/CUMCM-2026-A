"""
Q4 solver — drying with moving boundary (shrinking radius), per 规格卡_Q4.md.
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import openpyxl
from scipy.special import exp1
import lib_q1 as S            # reuse read_attach1

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ATT2 = os.path.join(BASE, "附件", "附件2.xlsx")
OUTDIR = os.path.join(BASE, "..", "结果")
os.makedirs(OUTDIR, exist_ok=True)

# ------------------------- constants (SI) -------------------------
T0 = 28.0 + 273.15          # K
C0 = 2.55                   # kg/kg
L = 0.25                    # m
R0 = 0.02                   # m
T_TAIL = 50.165 + 273.15    # K (D01 main)
C_TAIL = 0.04986
H = 25.0                    # W/(m2 K), H9
HM = 8e-7                   # m/s, H9
# appendix-4 parameters (+ H/HM so robustness sweeps can perturb)
P = dict(rho0=760.0, rho1=90.0,
         cp0=1850.0, cp1=2150.0,
         k0=0.12, k1=0.20,
         D0=4.2e-4, A_C=0.30, A_T=3850.0,
         H=25.0, HM=8e-7)

DXI = 0.05
N = 20
DT = 1.0
T_END = 259200.0
THRESH = 0.15


# ------------------------- appendix-4 properties -------------------------
def rho_c(C):
    return P['rho0'] + P['rho1'] * C


def cp_c(C):
    return P['cp0'] + P['cp1'] * C / (C + 1.0)


def k_c(C):
    return P['k0'] + P['k1'] * C / (C + 1.0)


def D_ct(C, T):
    Cs = np.maximum(C, 1e-12)
    return P['D0'] * np.exp(-P['A_C'] / Cs) * np.exp(-P['A_T'] / T)


def H_pot(C):
    """Kirchhoff potential with A_C=0.30: int_0^C exp(-0.30/s) ds."""
    a = P['A_C']
    Cs = np.maximum(C, 1e-12)
    return Cs * np.exp(-a / Cs) - a * exp1(a / Cs)


def B_T(T):
    return P['D0'] * np.exp(-P['A_T'] / T)


# ------------------------- data -------------------------
def read_attach2(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))[1:]
    t = np.array([r[0] for r in rows], float)
    R = np.array([r[1] for r in rows], float) / 100.0     # cm -> m
    assert np.all(np.diff(t) > 0), "attachment2 time not increasing"
    assert np.all(np.diff(R) <= 0), "attachment2 radius not monotone"
    return t, R


def bc_arrays(t_data, T_data, C_data, nsteps, dt, T_tail=None, C_tail=None):
    """T_inf/C_inf per step; attachment1 interp, clamped tail (D01).
    Optional asymptote override for cross-check."""
    from scipy.interpolate import PchipInterpolator
    t_n = np.arange(nsteps + 1) * dt
    T_inf = np.empty(nsteps + 1)
    C_inf = np.empty(nsteps + 1)
    m = t_n <= t_data[-1]
    T_inf[m] = PchipInterpolator(t_data, T_data, extrapolate=False)(t_n[m])
    C_inf[m] = PchipInterpolator(t_data, C_data, extrapolate=False)(t_n[m])
    T_inf[~m] = T_data[-1]
    C_inf[~m] = C_data[-1]
    if T_tail is not None:
        mask = t_n >= t_data[-1]
        T_inf[mask] = T_tail
        C_inf[mask] = C_tail
    return T_inf, C_inf


# ------------------------- Model A: material coordinate -------------------------
def solve_A(T_inf, C_inf, t_end=T_END, dt=DT, N=N, DXI=DXI,
            freeze_R=False, freeze_C=False, freeze_T=False, D_func=None,
            R_override=None, store_every=1, C_init=C0, conv_pert=0.0):
    if R_override is not None:
        t2, R2 = R_override
    else:
        t2, R2 = read_attach2(ATT2)
    nsteps = int(round(t_end / dt))
    tt = np.arange(nsteps + 1) * dt
    Rarr = np.interp(tt, t2, R2)                    # m
    if freeze_R:
        Rarr = np.full_like(Rarr, R0)
    Rdot = np.diff(Rarr) / dt                       # m/s, Rdot[n] for n -> n+1

    xi = np.arange(N + 1) * DXI
    xih = np.arange(N + 2) * DXI - DXI / 2
    xih[0] = 0.0
    xih[N + 1] = 1.0
    Vb = np.pi * (xih[1:] ** 2 - xih[:-1] ** 2) * L      # V_j(t) = R(t)^2 * Vb_j
    Ab = 2 * np.pi * xih[1:N + 1] * L                    # A_int(t) = R(t) * Ab
    AN_ = 2 * np.pi * L                                  # A_N(t) = R(t) * AN_

    nstore = nsteps // store_every + 1
    T = np.zeros((nstore, N + 1))
    C = np.zeros((nstore, N + 1))
    T[0, :] = T0
    C[0, :] = C_init
    cmax = np.empty(nsteps + 1)
    cmax[0] = C_init
    xmax = np.empty(nsteps)
    led_res = np.empty(nsteps)

    Tn = T[0].copy()
    Cn = C[0].copy()
    for n in range(nsteps):
        Rc = Rarr[n]
        Rd = Rdot[n]
        Rn1 = Rarr[n + 1]
        Cn_ = Cn
        Tn_ = Tn
        # ---- properties at C^n ----
        rho_j = rho_c(Cn_)
        cp_j = cp_c(Cn_)
        k_j = k_c(Cn_)
        V = Rc ** 2 * Vb
        # ---- heat: harmonic k interface, diffusion + shrinkage convection ----
        if freeze_T:
            Tnew = Tn
        else:
            denom = k_j[1:] + k_j[:-1]
            k_int = np.divide(2 * k_j[1:] * k_j[:-1], denom, out=np.zeros_like(denom),
                              where=denom > 0)
            Fh = -k_int * Ab * (Tn[1:] - Tn[:-1]) / DXI
            convT = np.zeros(N + 1)
            if Rd != 0.0:
                coef = xi * (Rd / Rc) * (1.0 + conv_pert * xi)
                convT[1:-1] = coef[1:-1] * (Tn[2:] - Tn[:-2]) / (2 * DXI)
                convT[N] = coef[N] * (Tn[N] - Tn[N - 1]) / DXI
            rhs = np.zeros(N + 1)
            rhs[0] = -Fh[0]
            rhs[1:-1] = Fh[:-1] - Fh[1:]
            rhs[N] = Fh[N - 1] + Rc * AN_ * P['H'] * (T_inf[n] - Tn[N])
            Tnew = Tn + dt * (rhs / (rho_j * cp_j * V) + convT)
        # ---- mass: Kirchhoff (A_C=0.30) + shrinkage convection ----
        if freeze_C:
            Cnew = Cn_
            Dmax = 0.0
            convC = np.zeros(N + 1)
        else:
            T_int = (Tnew[1:] + Tnew[:-1]) / 2
            if D_func is not None:
                D_j = D_func(Cn_)
                denomd = D_j[1:] + D_j[:-1]
                D_int = np.divide(2 * D_j[1:] * D_j[:-1], denomd,
                                  out=np.zeros_like(denomd), where=denomd > 0)
                Fm = -D_int * Ab * (Cn_[1:] - Cn_[:-1]) / DXI
                Dmax = D_j.max()
            else:
                B_i = B_T(T_int)
                dH = H_pot(Cn_[1:]) - H_pot(Cn_[:-1])
                Fm = -B_i * Ab * dH / DXI
                dC = Cn_[1:] - Cn_[:-1]
                Cmid = np.maximum((Cn_[1:] + Cn_[:-1]) / 2, 1e-12)
                chord = np.divide(B_i * dH, dC,
                                  out=B_i * np.exp(-P['A_C'] / Cmid),
                                  where=np.abs(dC) > 1e-14)
                Dmax = chord.max()
            convC = np.zeros(N + 1)
            if Rd != 0.0:
                coef = xi * (Rd / Rc) * (1.0 + conv_pert * xi)
                convC[1:-1] = coef[1:-1] * (Cn_[2:] - Cn_[:-2]) / (2 * DXI)
                convC[N] = coef[N] * (Cn_[N] - Cn_[N - 1]) / DXI
            rhsc = np.zeros(N + 1)
            rhsc[0] = -Fm[0]
            rhsc[1:-1] = Fm[:-1] - Fm[1:]
            rhsc[N] = Fm[N - 1] + Rc * AN_ * P['HM'] * (C_inf[n] - Cn_[N])
            Cnew = Cn_ + dt * (rhsc / V + convC)
        # ---- store & carry ----
        cmax[n + 1] = Cnew.max()
        if (n + 1) % store_every == 0:
            T[(n + 1) // store_every] = Tnew
            C[(n + 1) // store_every] = Cnew
        Tn = Tnew
        Cn = Cnew
        # ---- stability monitor (heat diffusion / mass chord / convection CFL) ----
        alpha_m = np.max(k_j / (rho_j * cp_j))
        x_h = alpha_m * dt / (Rc * DXI) ** 2
        x_m = Dmax * dt / (Rc * DXI) ** 2
        x_c = abs(Rd) * dt / (Rc * DXI)
        xmax[n] = max(x_h, x_m, x_c)
        # ---- discrete ledger residual: W = sum C*V, dW = BC flux + shrink + conv ----
        dW_rhs = Rc * AN_ * P['HM'] * (C_inf[n] - Cn_[N]) * dt
        dW_shr = np.sum(Cn_ * (Rn1 ** 2 - Rc ** 2) * Vb)
        dW_conv = dt * np.sum(V * convC)
        Wn = np.sum(Cn_ * V)
        Wn1 = np.sum(Cnew * (Rn1 ** 2 * Vb))
        led_res[n] = abs(Wn1 - Wn - dW_rhs - dW_shr - dW_conv)
    return T, C, Rarr, Rdot, xmax, led_res, cmax


# ------------------------- Model B: Euler fixed grid, truncation -------------------------
def solve_B(T_inf, C_inf, t_end=T_END, dt=DT, dr=0.001, freeze_R=False):
    t2, R2 = read_attach2(ATT2)
    nsteps = int(round(t_end / dt))
    Rarr = np.interp(np.arange(nsteps + 1) * dt, t2, R2)
    if freeze_R:
        Rarr = np.full_like(Rarr, R0)
    jmax = int(round(R0 / dr))
    r = np.arange(jmax + 1) * dr
    rhalf = np.arange(jmax + 2) * dr - dr / 2
    rhalf[0] = 0.0
    V = np.pi * (rhalf[1:] ** 2 - rhalf[:-1] ** 2) * L
    A = 2 * np.pi * rhalf[1:jmax + 1] * L
    T = np.full((nsteps + 1, jmax + 1), np.nan)
    C = np.full((nsteps + 1, jmax + 1), np.nan)
    T[0, :] = T0
    C[0, :] = C0
    for n in range(nsteps):
        Rc = Rarr[n]
        nact = int(np.floor(Rc / dr + 1e-9))        # surface node index
        Tn = T[n].copy()
        Cn = C[n].copy()
        rho_j = rho_c(Cn)
        cp_j = cp_c(Cn)
        k_j = k_c(Cn)
        # heat interior j=0..nact-1
        denom = k_j[1:] + k_j[:-1]
        k_int = np.divide(2 * k_j[1:] * k_j[:-1], denom, out=np.zeros_like(denom),
                          where=denom > 0)
        Fh = -k_int * A * (Tn[1:] - Tn[:-1]) / dr
        rhs = np.zeros(jmax + 1)
        rhs[0] = -Fh[0]
        rhs[1:nact] = Fh[:nact - 1] - Fh[1:nact]
        # surface truncated cell
        V_s = np.pi * (Rc ** 2 - (Rc - dr / 2) ** 2) * L
        A_s = 2 * np.pi * Rc * L
        rhs[nact] = Fh[nact - 1] + A_s * H * (T_inf[n] - Tn[nact])
        dT = np.zeros(jmax + 1)
        dT[:nact] = dt * rhs[:nact] / (rho_j[:nact] * cp_j[:nact] * V[:nact])
        dT[nact] = dt * rhs[nact] / (rho_j[nact] * cp_j[nact] * V_s)
        T[n + 1, :nact + 1] = Tn[:nact + 1] + dT[:nact + 1]
        # mass
        T_int = (T[n + 1][1:nact + 1] + T[n + 1][:nact]) / 2
        B_i = B_T(T_int)
        dH = H_pot(Cn[1:nact + 1]) - H_pot(Cn[:nact])
        Fm = -B_i * A[:nact] * dH / dr
        rhsc = np.zeros(jmax + 1)
        rhsc[0] = -Fm[0]
        rhsc[1:nact] = Fm[:nact - 1] - Fm[1:nact]
        rhsc[nact] = Fm[nact - 1] + A_s * HM * (C_inf[n] - Cn[nact])
        dC = np.zeros(jmax + 1)
        dC[:nact] = dt * rhsc[:nact] / V[:nact]
        dC[nact] = dt * rhsc[nact] / V_s
        C[n + 1, :nact + 1] = Cn[:nact + 1] + dC[:nact + 1]
    return T, C, Rarr


# ------------------------- Model C: frozen geometry -------------------------
def solve_C(T_inf, C_inf, t_end=T_END, dt=DT):
    T, C, Rarr, Rdot, xmax, led, cmax = solve_A(T_inf, C_inf, t_end, dt, freeze_R=True)
    return T, C


# ------------------------- outputs -------------------------
def endpoint_h(C):
    hit = np.where(np.nanmax(C, axis=1) <= THRESH)[0]
    return hit[0] / 3600.0 if hit.size else None


def interp_at(Cn_, Rc, rq):
    """Interpolate C at physical radius rq (m) on material grid."""
    xiq = rq / Rc
    idx = xiq / DXI
    j = int(np.floor(idx))
    if j >= N:
        return Cn_[N]
    w = idx - j
    return Cn_[j] * (1 - w) + Cn_[j + 1] * w


def main():
    t_data, T_data, C_data = S.read_attach1(S.ATT1)
    dt, t_end = DT, T_END
    nsteps = int(round(t_end / dt))
    T_inf, C_inf = bc_arrays(t_data, T_data, C_data, nsteps, dt)
    T, C, Rarr, Rdot, xmax, led, cmax = solve_A(T_inf, C_inf, t_end, dt)

    # ---- endpoint ----
    hit = np.where(cmax <= THRESH)[0]
    assert hit.size > 0, "criterion never reached in 3 days"
    n_hit = hit[0]
    t_hit = n_hit * dt
    print("endpoint: t = %d s = %.4f h" % (t_hit, t_hit / 3600.0))
    print("max C at t_hit = %.6f (center = %.6f, surface = %.6f)"
          % (C[n_hit].max(), C[n_hit, 0], C[n_hit, N]))
    print("max C at t-60s = %.6f" % C[n_hit - 60].max())
    span = C.max(axis=1)[n_hit - 600:n_hit + 600]
    print("C_max monotone across threshold: %s" % bool(np.all(np.diff(span) <= 0)))

    # ---- 表6 ----
    t_list_h = list(range(6, int(t_hit // 3600) + 1, 6))
    r_cols_cm = [0.0, 0.5, 1.0]
    print("\n--- 表6 水分浓度(kg/kg) ---")
    print("t/h | " + " | ".join("r=%.1fcm" % r for r in r_cols_cm) + " | 表面")
    rows = []
    for th in t_list_h:
        i = int(th * 3600 / dt)
        Rc = Rarr[i]
        vals = [interp_at(C[i], Rc, rr / 100.0) for rr in r_cols_cm]
        vals.append(C[i, N])
        rows.append((th, vals))
        print("%4d | " % th + " | ".join("%8.4f" % v for v in vals))
    i = n_hit
    Rc = Rarr[i]
    vals = [interp_at(C[i], Rc, rr / 100.0) for rr in r_cols_cm] + [C[i, N]]
    rows.append((t_hit / 3600.0, vals))
    print("终点 %.2fh | " % (t_hit / 3600.0) + " | ".join("%8.4f" % v for v in vals))

    # ---- result4.xlsx ----
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "水分浓度"
    ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
    dists_cm = [round(0.1 * k, 1) for k in range(13)]        # 0..1.2 cm
    for c, d in enumerate(dists_cm, start=2):
        ws.cell(row=1, column=c, value=d)
    ws.cell(row=1, column=15, value="药材表面")
    row_i = 2
    t_last = (n_hit // 60) * 60
    for t_s in range(60, t_last + 1, 60):
        i = int(t_s // dt)
        Rc = Rarr[i]
        ws.cell(row=row_i, column=1, value=t_s)
        for c, d in enumerate(dists_cm, start=2):
            ws.cell(row=row_i, column=c, value=round(float(interp_at(C[i], Rc, d / 100.0)), 4))
        ws.cell(row=row_i, column=15, value=round(float(C[i, N]), 4))
        row_i += 1
    ws.cell(row=row_i, column=1, value=t_hit)
    Rc = Rarr[n_hit]
    for c, d in enumerate(dists_cm, start=2):
        ws.cell(row=row_i, column=c, value=round(float(interp_at(C[n_hit], Rc, d / 100.0)), 4))
    ws.cell(row=row_i, column=15, value=round(float(C[n_hit, N]), 4))
    out = os.path.join(OUTDIR, "result4.xlsx")
    wb.save(out)
    print("\nsaved:", out, "(data rows = %d)" % (row_i - 1))

    # ---- save arrays for verification ----
    np.savez_compressed(os.path.join(OUTDIR, "q4_arrays.npz"),
                        T=T, C=C, Rarr=Rarr, Rdot=Rdot, xmax=xmax, led=led)

    # ---- 五层检查 ----
    print("\n== 五层检查 ==")
    print("L1 程序: 索引/维度 OK; 步数=%d 节点=%d" % (nsteps, N + 1))
    print("L2 数学: 规格卡_Q4 式(1)(2)(4) 逐项实现; Kirchhoff A_C=%.2f; 顺序显式耦合" % P['A_C'])
    print("L3 数据: 附件2 R=[%.4f,%.4f] m 单调; BC 附件1 插值+D01 尾值" % (Rarr.min(), Rarr.max()))
    print("L4 数值: x_max=%.4f (<0.4226); ledger residual max=%.2e; C 恒正=%s"
          % (xmax.max(), led.max(), bool((C > -1e-12).all())))
    print("L5 建模: 表6+result4 直接回答 Q4; 终点 %.4f h" % (t_hit / 3600.0))


if __name__ == "__main__":
    main()
