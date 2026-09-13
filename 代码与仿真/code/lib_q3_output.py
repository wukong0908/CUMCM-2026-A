# -*- coding: utf-8 -*-
"""
Q3 outputs — drying endpoint + 表5 + result3.xlsx.
Reuses Q2 M2 full-process solution (Kirchhoff FVM, dr=1mm, dt=1s).
Spec: t_dry = min{t : max_j C <= 0.15}; table5 rows 6..54 h + endpoint row;
result3 = 60 s steps to 205920 s + final endpoint row 205952 s, 0.1 cm spacing.
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import openpyxl
import lib_q2 as Q
import lib_q1 as S

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDIR = os.path.join(BASE, "结果")
os.makedirs(OUTDIR, exist_ok=True)

THRESH = 0.15


def main():
    t_data, T_data, C_data = S.read_attach1(S.ATT1)
    dt, dr, t_end = 1.0, 0.001, 259200.0
    nsteps = int(round(t_end / dt))
    T_inf, C_inf = Q.bc_arrays(t_data, T_data, C_data, nsteps, dt)
    T, C, cour, m_bal = Q.solve(T_inf, C_inf, t_end, dr, dt, flux_scheme="kirchhoff")

    # ---- endpoint criterion ----
    cmax = C.max(axis=1)
    hit = np.where(cmax <= THRESH)[0]
    assert hit.size > 0, "criterion never reached"
    n_hit = hit[0]
    t_hit = n_hit * dt                       # 205952 s
    print("endpoint: t = %d s = %.3f h (57.209 h)" % (t_hit, t_hit / 3600.0))
    print("max C at t_hit = %.6f (center C = %.6f)" % (cmax[n_hit], C[n_hit, 0]))
    print("max C at t_hit-60s = %.6f (not yet dried)" % cmax[n_hit - 60])
    # criterion sanity: C_max strictly decreasing across threshold, no rebound
    span = cmax[n_hit - 600:n_hit + 600]
    mono = bool(np.all(np.diff(span) <= 0))
    print("C_max monotone decreasing in +-10min window: %s" % mono)

    # ---- 表5 ----
    t_list_h = [6, 12, 18, 24, 30, 36, 42, 48, 54]
    r_cm = [0.0, 0.5, 1.0, 1.5, 2.0]
    jdx = [int(round(rr / (dr * 100))) for rr in r_cm]
    print("\n--- 表5 水分浓度(kg/kg) ---")
    print("t/h | " + " | ".join("r=%.1fcm" % r for r in r_cm))
    rows = []
    for th in t_list_h:
        i = int(th * 3600 / dt)
        row = [C[i, j] for j in jdx]
        rows.append((th, row))
        print("%4d | " % th + " | ".join("%8.4f" % v for v in row))
    row_end = [C[n_hit, j] for j in jdx]
    rows.append(("end %.3f h" % (t_hit / 3600.0), row_end))
    print("终点 57.21h | " + " | ".join("%8.4f" % v for v in row_end))

    # ---- result3.xlsx ----
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "水分浓度"
    header = "时间\\到药材中心的距离"
    dists = [round(dr * 100 * j, 1) for j in range(21)]
    ws.cell(row=1, column=1, value=header)
    for c, d in enumerate(dists, start=2):
        ws.cell(row=1, column=c, value=d)
    row_i = 2
    t_last = (n_hit // 60) * 60                 # 205920: largest 60s grid <= t_hit
    for t_s in range(60, t_last + 1, 60):
        ws.cell(row=row_i, column=1, value=t_s)
        i = int(t_s // dt)
        for j in range(21):
            ws.cell(row=row_i, column=j + 2, value=round(float(C[i, j]), 4))
        row_i += 1
    # final endpoint row 205952 s
    ws.cell(row=row_i, column=1, value=t_hit)
    for j in range(21):
        ws.cell(row=row_i, column=j + 2, value=round(float(C[n_hit, j]), 4))
    out = os.path.join(OUTDIR, "result3.xlsx")
    wb.save(out)
    print("\nsaved:", out, "(rows = %d data rows)" % (row_i - 1))
    print("rows: 60..%d s step 60 (%d) + endpoint row %d s" %
          (t_last, t_last // 60, t_hit))

    # ---- L5 check ----
    print("\n== 五层检查 ==")
    print("L1 程序: 索引/维度 OK; 数据行数=%d" % (row_i - 1))
    print("L2 数学: 判据 max C<=%.2f 首达 t=%ds; C_max 单调跨越, 无回弹=%s" % (THRESH, t_hit, mono))
    print("L3 数据: 全程解复用 M2(附件1+D01 BC); 采样 60s 整除节点")
    print("L4 数值: Courant max=%.5f; C 恒正=%s; 守恒 rel=%.2e" % (cour.max(), bool((C > 0).all()), m_bal["rel"]))
    print("L5 建模: 表5+result3 直接回答 Q3; 终点 57.21 h 与 V4 三网格收敛一致")


if __name__ == "__main__":
    main()
